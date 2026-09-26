from pathlib import Path
from typing import List, Dict
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv"}


def extract_sections(path: Path) -> List[Dict]:
    """Return document content with location metadata preserved."""
    suffix = path.suffix.lower()
    sections = []

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        for page_no, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.replace("\x00", " ").strip()
            if text:
                sections.append({"text": text, "location": f"page {page_no}"})
        return sections

    if suffix == ".docx":
        doc = Document(str(path))
        for i, paragraph in enumerate(doc.paragraphs, start=1):
            if paragraph.text.strip():
                sections.append({"text": paragraph.text.strip(), "location": f"paragraph {i}"})
        for ti, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                vals = [cell.text.strip() for cell in row.cells]
                if any(vals):
                    rows.append(" | ".join(vals))
            if rows:
                sections.append({"text": "\n".join(rows), "location": f"table {ti}"})
        return sections

    if suffix == ".pptx":
        prs = Presentation(str(path))
        for slide_no, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            if parts:
                sections.append({"text": "\n".join(parts), "location": f"slide {slide_no}"})
        return sections

    if suffix == ".xlsx":
        wb = load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sections.append({"text": "\n".join(rows), "location": f"sheet {ws.title}"})
        return sections

    if suffix in {".txt", ".md", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        if text:
            sections.append({"text": text, "location": "text file"})
        return sections

    raise ValueError(f"Unsupported file type: {suffix}")


def extract_text(path: Path) -> str:
    return "\n\n".join(section["text"] for section in extract_sections(path))
