from pathlib import Path
from typing import List, Dict
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def _safe_title(title: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in " -_" else "_" for c in title)
    return cleaned.strip()[:80] or "generated_document"


def make_docx(title: str, content: str, output: Path):
    doc = Document()
    doc.add_heading(title, 0)
    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            doc.add_paragraph(paragraph)
    doc.save(output)
    return output


def make_pptx(title: str, content: str, output: Path):
    prs = Presentation()
    chunks = [p.strip() for p in content.split("\n\n") if p.strip()]
    if not chunks:
        chunks = [content]

    for idx, chunk in enumerate(chunks[:20]):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title if idx == 0 else f"{title} — {idx + 1}"
        slide.placeholders[1].text = chunk[:4000]

    prs.save(output)
    return output


def make_xlsx(title: str, content: str, output: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "AI Response"
    ws["A1"] = title
    ws["A1"].font = ws["A1"].font.copy(bold=True)
    row = 3
    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            ws.cell(row=row, column=1, value=paragraph)
            row += 1
    ws.column_dimensions["A"].width = 110
    wb.save(output)
    return output


def make_pdf(title: str, content: str, output: Path):
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(output), pagesize=A4)
    story = [Paragraph(title, styles["Title"]), Spacer(1, 16)]

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            safe = (
                paragraph.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            story.append(Paragraph(safe, styles["BodyText"]))
            story.append(Spacer(1, 8))

    doc.build(story)
    return output


def generate(kind: str, title: str, content: str, output_dir: Path):
    base = _safe_title(title)

    if kind == "docx":
        path = output_dir / f"{base}.docx"
        return make_docx(title, content, path)

    if kind == "pptx":
        path = output_dir / f"{base}.pptx"
        return make_pptx(title, content, path)

    if kind == "xlsx":
        path = output_dir / f"{base}.xlsx"
        return make_xlsx(title, content, path)

    if kind == "pdf":
        path = output_dir / f"{base}.pdf"
        return make_pdf(title, content, path)

    raise ValueError("Unsupported output format")
